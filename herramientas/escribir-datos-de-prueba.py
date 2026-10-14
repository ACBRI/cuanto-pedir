#!/usr/bin/env python3
"""Escribe los archivos que usa la prueba de punta a punta (pruebas/datos/).

- consumo-windows-1252.csv: consumo diario de materias primas en kilos, como lo
  guarda Excel en Windows: punto y coma, coma decimal, fechas día/mes/año y
  codificación Windows-1252 (no UTF-8).
- ventas.xlsx: un libro de Excel con una hoja de notas y una de ventas que
  empieza con una fila de título; las fechas son fechas de Excel, no texto.

Los datos son inventados, con semilla fija.

  python3 herramientas/escribir-datos-de-prueba.py
"""
from __future__ import annotations

import datetime as dt
import random
from pathlib import Path

from openpyxl import Workbook

CARPETA = Path(__file__).resolve().parent.parent / "pruebas" / "datos"


def consumo() -> None:
    azar = random.Random(7)
    materias = [
        ("MP-101", "Lactosa monohidrato", 12.0),
        ("MP-102", "Almidón de maíz", 6.5),
        ("MP-103", "Estearato de magnesio", 0.8),
        ("MP-104", "Celulosa microcristalina", 9.25),
    ]
    inicio = dt.date(2026, 6, 1)
    lineas = ["Fecha;Código;Descripción;Cantidad (kg)"]
    for k in range(120):
        dia = inicio + dt.timedelta(days=k)
        for codigo, nombre, media in materias:
            if codigo == "MP-103" and azar.random() < 0.6:
                continue
            kilos = max(0.05, azar.gauss(media, media * 0.3))
            texto = f"{kilos:.2f}".replace(".", ",")
            lineas.append(f"{dia.day}/{dia.month}/{dia.year};{codigo};{nombre};{texto}")
    (CARPETA / "consumo-windows-1252.csv").write_bytes(("\r\n".join(lineas) + "\r\n").encode("cp1252"))


def ventas() -> None:
    azar = random.Random(11)
    libro = Workbook()
    notas = libro.active
    notas.title = "Notas"
    notas.append(["Este libro es de prueba: la hoja de ventas es la segunda."])
    hoja = libro.create_sheet("Ventas")
    hoja.append(["Informe de ventas por artículo"])
    hoja.append([])
    hoja.append(["Fecha", "Artículo", "Unidades", "Precio unitario"])
    inicio = dt.datetime(2026, 5, 4)
    for k in range(150):
        dia = inicio + dt.timedelta(days=k)
        if dia.weekday() == 6:
            continue
        hoja.append([dia, "A-1", azar.randint(3, 9), 2.5])
        if azar.random() < 0.15:
            hoja.append([dia, "B-2", azar.randint(1, 4), 11.0])
    for fila in hoja.iter_rows(min_row=4, max_col=1):
        fila[0].number_format = "dd/mm/yyyy"
    libro.save(CARPETA / "ventas.xlsx")


if __name__ == "__main__":
    CARPETA.mkdir(parents=True, exist_ok=True)
    consumo()
    ventas()
    for archivo in sorted(CARPETA.iterdir()):
        print(archivo.relative_to(CARPETA.parent.parent), archivo.stat().st_size, "bytes")
